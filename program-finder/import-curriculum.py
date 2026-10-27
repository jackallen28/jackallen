#!/usr/bin/env python3
"""Regenerate the curriculum data block in program-finder.html from the curriculum links workbook.

Usage: python3 -I program-finder/import-curriculum.py Circuit_Pathways_Curriculum_Links.xlsx

Reads the 'Levels 7-8', 'Levels 9-10' and 'VCE' tabs and rewrites the JSON between the
curriculum-data markers in the page. Dev-only tooling; nothing here ships to teachers.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

HTML = Path(__file__).with_name('program-finder.html')
START, END = '<!-- curriculum-data:start -->', '<!-- curriculum-data:end -->'


def platforms(text):
    """Normalise the 'Best platform(s)' cell to tokens: microbit, arduino, esp, pi, all."""
    out = set()
    for tok in re.split(r'[,;]', text or ''):
        t = re.sub(r'\(.*?\)', '', tok).strip().lower()
        if not t:
            continue
        if t in ('all', 'any') or t.startswith('all'):
            out.add('all')
        elif 'micro:bit' in t:
            out.add('microbit')
        elif 'raspberry' in t:
            out.add('pi')
        elif 'wifi' in t or t == 'esp':
            out.add('esp')
        elif 'arduino' in t:
            out.add('arduino')
    return sorted(out) or ['all']


def strength(text):
    t = (text or '').lower()
    kind = 'core' if t.startswith('core') else 'supporting' if t.startswith('supporting') else 'extension'
    only = None
    if '(car)' in t:
        only = 'car'
    elif 'speed camera' in t:
        only = 'speed'
    elif '(pi)' in t:
        only = 'pi'
    return kind, only


def clean(v):
    if v is None:
        return None
    # Arrows can render as emoji on some devices, so spell them out.
    text = str(v).replace('\u2194', 'and').replace('\u2192', 'to')
    return re.sub(r'\s+', ' ', text).strip()


def f10_rows(ws, band):
    rows = []
    for r in ws.iter_rows(min_row=1, values_only=True):
        area, strand, code, text, binlink, carlink, plat, strg = (clean(c) for c in r[:8])
        if not code or not re.match(r'^VC2', code):
            continue
        kind, only = strength(strg)
        rows.append({
            'band': band, 'area': area, 'strand': strand, 'code': code, 'text': text,
            'bin': None if binlink in (None, '—') else binlink,
            'car': None if carlink in (None, '—') else carlink,
            'platforms': platforms(plat), 'strength': kind, 'only': only,
            'confirm': code == 'VC2TDI10C04',  # flagged on the Overview tab
        })
    return rows


def vce_rows(ws):
    rows = []
    for r in ws.iter_rows(min_row=1, values_only=True):
        a, unit, typ, text, binlink, carlink, plat, strg = (clean(c) for c in r[:8])
        if a in ('Systems Engineering', 'Applied Computing', 'Algorithmics') and unit and text:
            kind, only = strength(strg)
            rows.append({
                'band': 'vce', 'study': a, 'unit': unit, 'type': typ, 'text': text,
                'bin': None if binlink in (None, '—') else binlink,
                'car': None if carlink in (None, '—') else carlink,
                'platforms': platforms(plat), 'strength': kind, 'only': only,
            })
    return rows


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    data = f10_rows(wb['Levels 7-8'], '7-8') + f10_rows(wb['Levels 9-10'], '9-10') + vce_rows(wb['VCE'])
    block = json.dumps(data, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    html = HTML.read_text(encoding='utf-8')
    pat = re.compile(re.escape(START) + r'.*?' + re.escape(END), re.S)
    if not pat.search(html):
        sys.exit('curriculum-data markers not found in program-finder.html')
    new = START + '\n<script type="application/json" id="curriculum-data">\n' + block + '\n</script>\n' + END
    HTML.write_text(pat.sub(lambda m: new, html), encoding='utf-8')
    counts = {b: sum(1 for d in data if d['band'] == b) for b in ('7-8', '9-10', 'vce')}
    print('Wrote', len(data), 'rows', counts)


if __name__ == '__main__':
    main()
