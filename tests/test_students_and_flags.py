"""Students, serial numbers and flags, through the API the pages use.

A named student gets their sheets logged, is not given the same question
twice unless asked, and has workbooks rewritten after every sheet. A serial
is printed on every question and finds it again. A flagged question never
goes on a sheet until the flag is cleared.
"""

import pytest

from blitz.corpus.sample import load_all_samples
from blitz import db

from tests.test_indexer_ui import client as _client  # noqa: F401  (fixture)


@pytest.fixture
def seeded_client(_client, tmp_path):
    with db.session(tmp_path / "t.sqlite3") as conn:
        load_all_samples(conn)
    return _client


def _kk(client, subject="physics"):
    data = client.get("/api/subjects").json()
    s = next(x for x in data["subjects"] if x["id"] == subject)
    return [k["id"] for u in s["units"] for a in u["areas"] for k in a["key_knowledge"]
            if k["available"]]


def test_every_question_has_a_serial_and_it_finds_the_question(seeded_client):
    res = seeded_client.get("/api/questions", params={"subject": "physics", "q": ""})
    qs = res.json()["questions"]
    assert qs and all(q["serial"] and q["serial"].startswith("PH-") for q in qs)
    serial = qs[3]["serial"]
    hit = seeded_client.get("/api/questions", params={"subject": "physics", "q": serial}).json()
    assert hit["total"] == 1 and hit["questions"][0]["serial"] == serial
    # Words find it too, and a lower-case serial is fine.
    word = next(w for w in qs[3]["body"].split() if len(w) > 5 and w.isalpha())
    words = seeded_client.get("/api/questions", params={"subject": "physics", "q": word}).json()
    assert any(q["serial"] == serial for q in words["questions"])
    assert seeded_client.get("/api/questions", params={
        "subject": "physics", "q": serial.lower()}).json()["total"] == 1


def test_a_named_student_is_logged_and_not_given_repeats(seeded_client, tmp_path):
    # The sample bank is small: a few dot points first, so there is something
    # left for the second sheet to draw on.
    kk = _kk(seeded_client)
    few, many = kk[:3], kk
    res = seeded_client.post("/api/generate", json={
        "subject_id": "physics", "kk_ids": few, "seed": 1,
        "student_name": "Year 12 Physics"})
    assert res.status_code == 200, res.text
    first = res.json()
    assert first["student"]["name"] == "Year 12 Physics"
    assert first["serials"] and first["skipped"] == []
    sid = first["student"]["id"]

    students = seeded_client.get("/api/students").json()["students"]
    me = next(s for s in students if s["id"] == sid)
    assert me["sheets"] == 1 and me["questions"] == len(first["serials"])

    res = seeded_client.post("/api/generate", json={
        "subject_id": "physics", "kk_ids": many, "seed": 1, "student_id": sid})
    assert res.status_code == 200, res.text
    second = res.json()
    assert set(second["serials"]).isdisjoint(first["serials"])
    assert set(second["skipped"]) >= set(first["serials"])
    assert any("already had" in w for w in second["warnings"])

    # Allow repeats: the same request as the first picks the same questions.
    res = seeded_client.post("/api/generate", json={
        "subject_id": "physics", "kk_ids": few, "seed": 1, "student_id": sid,
        "allow_repeats": True})
    assert res.status_code == 200, res.text
    again = res.json()
    assert again["skipped"] == []
    assert set(again["serials"]) == set(first["serials"])

    detail = seeded_client.get(f"/api/students/{sid}").json()
    assert len(detail["sheets"]) == 3
    assert detail["sheets"][0]["url"].startswith("/sheets/")
    assert all(q["serial"] for sh in detail["sheets"] for q in sh["questions"])

    # Workbooks: the master and one for the student, both readable.
    from openpyxl import load_workbook

    master = load_workbook(tmp_path / "students" / "MASTER.xlsx")
    assert set(master.sheetnames) >= {"Students", "Log", "Questions"}
    log = master["Log"]
    rows = list(log.iter_rows(min_row=2, values_only=True))
    assert len(rows) == me["questions"] + len(second["serials"]) + len(again["serials"])
    assert all(r[1] == "Year 12 Physics" for r in rows)
    mine = load_workbook(tmp_path / "students" / "Year 12 Physics.xlsx")
    assert "Questions given" in mine.sheetnames
    assert seeded_client.get("/students/files/Year%2012%20Physics.xlsx").status_code == 200


def test_the_serial_is_printed_on_the_sheet(seeded_client, tmp_path):
    import pymupdf

    kk = _kk(seeded_client)[:4]
    res = seeded_client.post("/api/generate", json={
        "subject_id": "physics", "kk_ids": kk, "seed": 2}).json()
    pdf = tmp_path / "sheets" / res["filename"]
    text = " ".join(page.get_text() for page in pymupdf.open(pdf))
    for serial in res["serials"][:3]:
        assert serial in text


def test_a_flagged_question_is_never_picked_until_cleared(seeded_client):
    kk = _kk(seeded_client)[:6]
    plan = seeded_client.post("/api/preview", json={
        "subject_id": "physics", "kk_ids": kk, "seed": 3}).json()
    victim = plan["questions"][0]
    res = seeded_client.post(f"/api/questions/{victim['id']}/flag",
                             json={"flag": "incomplete", "note": "options cut off"})
    assert res.status_code == 200 and res.json()["flag"] == "incomplete"

    after = seeded_client.post("/api/preview", json={
        "subject_id": "physics", "kk_ids": kk, "seed": 3}).json()
    assert victim["id"] not in {q["id"] for q in after["questions"]}

    flagged = seeded_client.get("/api/questions", params={
        "subject": "physics", "flagged": "1"}).json()
    assert [q["id"] for q in flagged["questions"]] == [victim["id"]]
    assert flagged["questions"][0]["flag_note"] == "options cut off"

    assert seeded_client.post(f"/api/questions/{victim['id']}/flag",
                              json={"flag": "nonsense"}).status_code == 400
    seeded_client.post(f"/api/questions/{victim['id']}/flag", json={"flag": None})
    back = seeded_client.post("/api/preview", json={
        "subject_id": "physics", "kk_ids": kk, "seed": 3}).json()
    assert victim["id"] in {q["id"] for q in back["questions"]}


def test_serials_and_flags_survive_a_re_import(tmp_path):
    from blitz.db import insert_question, set_flag, upsert_source

    with db.session(tmp_path / "x.sqlite3") as conn:
        upsert_source(conn, id="src", subject_id="physics", kind="other", title="Src")
        row = {"id": "src-q1", "subject_id": "physics", "source_id": "src",
               "question_type": "ph-mc", "body": "First version"}
        insert_question(conn, row, ["physics-u3-aos1-kk01"])
        serial = conn.execute("SELECT serial FROM question").fetchone()["serial"]
        assert serial == "PH-0001"
        set_flag(conn, "src-q1", "wrong", "answer key disagrees")
        insert_question(conn, {**row, "body": "Second version"}, ["physics-u3-aos1-kk01"])
        got = conn.execute("SELECT serial, flag, flag_note, body FROM question").fetchone()
        assert (got["serial"], got["flag"], got["body"]) == (serial, "wrong", "Second version")
        insert_question(conn, {**row, "id": "src-q2"}, ["physics-u3-aos1-kk01"])
        assert conn.execute("SELECT serial FROM question WHERE id='src-q2'").fetchone()[0] == "PH-0002"
