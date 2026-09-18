"""Word files in: worksheets and notes as .docx, and VCAA's Word study designs.

A .docx is rendered to a PDF with the typography the extractors expect and then
indexed like any other book; a study design is read as text with its dot
points bulleted. These tests build small Word files with python-docx so they
never depend on a licensed document.
"""

import io

import pytest

docx = pytest.importorskip("docx")

from docx.shared import Inches  # noqa: E402


def _worksheet(path):
    """A teacher's revision worksheet: headings, prose, numbered questions,
    an image and a table. Numbering comes from Word's own 'List Number'
    style, which carries it on the style rather than the paragraph."""
    from PIL import Image, ImageDraw

    d = docx.Document()
    d.add_heading("Unit 3 Motion revision", level=1)
    d.add_heading("Projectile motion", level=2)
    p = d.add_paragraph("A projectile launched horizontally keeps a constant "
                        "horizontal velocity while its vertical velocity "
                        "increases at 9.8 m s")
    p.add_run("-2").font.superscript = True
    p.add_run(". The range depends on the launch height and the horizontal speed.")
    d.add_paragraph("Questions", style="Heading 3")
    q = d.add_paragraph("A ball is thrown horizontally at 12 m s", style="List Number")
    q.add_run("-1").font.superscript = True
    q.add_run(" from a cliff 20 m high. How far from the base does it land? (3 marks)")
    d.add_paragraph("A projectile is launched at 30° to the horizontal at 25 m s⁻¹. "
                    "Which of the following is closest to its time of flight? (1 mark)",
                    style="List Number")
    for opt in ("A. 1.3 s", "B. 2.6 s", "C. 3.9 s", "D. 5.1 s"):
        d.add_paragraph(opt)
    d.add_paragraph("Explain why the horizontal component of velocity is unchanged "
                    "in flight. (2 marks)", style="List Number")
    im = Image.new("RGB", (400, 200), (255, 255, 255))
    ImageDraw.Draw(im).line((10, 190, 390, 20), fill=(0, 0, 0), width=4)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    d.add_picture(buf, width=Inches(3))
    d.add_heading("Circular motion", level=2)
    d.add_paragraph("An object in uniform circular motion has a centripetal "
                    "acceleration directed towards the centre of the circle.")
    t = d.add_table(rows=2, cols=3)
    for j, h in enumerate(("radius (m)", "speed (m/s)", "period (s)")):
        t.cell(0, j).text = h
    for j, v in enumerate(("2.0", "4.0", "3.1")):
        t.cell(1, j).text = v
    d.add_paragraph("A 2.0 kg mass moves in a circle of radius 1.5 m at 3.0 m s⁻¹. "
                    "Calculate the centripetal force. (2 marks)", style="List Number")
    d.save(path)
    return path


def _study_design(path):
    """The shape of VCAA's Word study design: unit and area headings, an
    outcome, then 'Key knowledge' and 'Key skills' as bulleted lists."""
    d = docx.Document()
    d.add_heading("Unit 3: Managing a business", level=1)
    d.add_paragraph("In this unit students explore the key processes for managing a business.")
    d.add_heading("Area of Study 1", level=2)
    d.add_heading("Business foundations", level=3)
    d.add_paragraph("This area of study introduces students to the characteristics of businesses.")
    d.add_heading("Outcome 1", level=3)
    d.add_paragraph("On completion of this unit the student should be able to analyse "
                    "the key characteristics of businesses and their stakeholders.")
    d.add_paragraph("To achieve this outcome the student will draw on key knowledge "
                    "and key skills outlined in Area of Study 1.")
    d.add_heading("Key knowledge", level=5)
    for kk in ("types of businesses including sole traders, partnerships and companies",
               "business objectives including to make a profit and to increase market share",
               "stakeholders of businesses including owners, managers and employees"):
        d.add_paragraph(kk, style="List Bullet")
    d.add_heading("Key skills", level=5)
    d.add_paragraph("identify, define, describe and apply business management concepts",
                    style="List Bullet")
    d.add_heading("Area of Study 2", level=2)
    d.add_heading("Human resource management", level=3)
    d.add_heading("Outcome 2", level=3)
    d.add_paragraph("On completion of this unit the student should be able to explain "
                    "theories of motivation.")
    d.add_heading("Key knowledge", level=5)
    for kk in ("the relationship between human resource management and business objectives",
               "key principles of the theories of motivation of Maslow, Locke and Lawrence"):
        d.add_paragraph(kk, style="List Bullet")
    d.add_heading("Key skills", level=5)
    d.add_paragraph("analyse case studies of business management", style="List Bullet")
    d.add_heading("Unit 4: Transforming a business", level=1)
    d.add_heading("Area of Study 1", level=2)
    d.add_heading("Reviewing performance", level=3)
    d.add_heading("Outcome 1", level=3)
    d.add_paragraph("On completion of this unit the student should be able to explain "
                    "the way business change may come about.")
    d.add_heading("Key knowledge", level=5)
    for kk in ("the concept of business change and its causes",
               "proactive and reactive approaches to change in a business"):
        d.add_paragraph(kk, style="List Bullet")
    d.add_heading("Key skills", level=5)
    d.add_paragraph("interpret and evaluate business information", style="List Bullet")
    d.save(path)
    return path


def test_word_numbering_becomes_visible_and_bullets_are_bulleted(tmp_path):
    from blitz.ingest.docx import docx_text

    text = docx_text(_worksheet(tmp_path / "ws.docx"))
    lines = [ln for ln in text.splitlines() if ln.strip()]
    import re

    numbered = [ln for ln in lines if re.match(r"[1-4]\. [A-Z]", ln)]
    assert len(numbered) == 4, lines
    assert numbered[0].startswith("1. A ball is thrown")
    assert numbered[3].startswith("4. A 2.0 kg mass")
    # A table survives as rows.
    assert "radius (m) | speed (m/s) | period (s)" in lines


def test_word_file_renders_to_a_pdf_the_extractors_can_read(tmp_path):
    import pymupdf

    from blitz.ingest.docx import docx_to_pdf
    from blitz.ingest.textflow import build_lines

    pdf = docx_to_pdf(_worksheet(tmp_path / "ws.docx"), out_dir=tmp_path)
    assert pdf.exists() and pdf.suffix == ".pdf"
    doc = pymupdf.open(pdf)
    lines = build_lines(doc[0])
    by_text = {ln.text[:22]: ln for ln in lines}
    # Headings are set larger than the body, which is how passages find them.
    assert by_text["Unit 3 Motion revision"].size > by_text["Projectile motion"].size
    body = next(ln for ln in lines if ln.text.startswith("A projectile launched"))
    assert by_text["Projectile motion"].size > body.size
    # Superscripts, whether set in Word or typed as Unicode, come back as text.
    joined = " ".join(ln.text for ln in lines)
    assert "9.8 m s⁻²" in joined
    assert "12 m s⁻¹" in joined and "25 m s⁻¹" in joined
    assert "1. A ball is thrown" in joined
    assert len(doc[0].get_images()) == 1


def test_a_word_worksheet_indexes_as_a_book(tmp_path, monkeypatch):
    from blitz import config, db
    from blitz.ingest.index import index_book

    monkeypatch.setattr(config, "CROPS_DIR", tmp_path / "crops")
    import blitz.ingest.extract as extract

    monkeypatch.setattr(extract, "CROPS_DIR", tmp_path / "crops", raising=False)
    path = _worksheet(tmp_path / "motion-worksheet.docx")
    with db.session(tmp_path / "t.sqlite3") as conn:
        report = index_book(conn, path, subject_id="physics", source_id="ws",
                            progress=lambda *a: None)
        rows = conn.execute(
            "SELECT question_type, marks, body FROM question ORDER BY id").fetchall()
        passages = conn.execute("SELECT title FROM passage").fetchall()
    assert report.kind == "textbook"
    assert report.questions_indexed >= 3
    assert any("converted from Word" in r for r in report.review)
    types = {r["question_type"] for r in rows}
    assert "ph-mc" in types
    marks = {r["marks"] for r in rows}
    assert {1, 2, 3} <= marks
    assert any("m s⁻¹" in r["body"] for r in rows)
    assert passages, "the teaching prose under a heading should be indexed"


def test_a_word_study_design_imports(tmp_path):
    import yaml

    from blitz.studydesign.importer import import_study_design

    out = import_study_design(_study_design(tmp_path / "bm-sd.docx"),
                              "business-management", out_dir=tmp_path)
    data = yaml.safe_load(out.read_text(encoding="utf-8"))
    units = data["units"]
    assert [u["number"] for u in units] == [3, 4]
    u3 = units[0]
    assert u3["title"] == "Managing a business"
    areas = u3["areas_of_study"]
    assert [a["title"] for a in areas] == ["Business foundations",
                                           "Human resource management"]
    assert areas[0]["outcome"].startswith("On completion of this unit")
    assert "To achieve this outcome" not in areas[0]["outcome"]
    kk = [k["text"] for k in areas[0]["key_knowledge"]]
    assert len(kk) == 3 and kk[0].startswith("types of businesses")
    assert all(k["verified"] for a in areas for k in a["key_knowledge"])
    assert units[1]["areas_of_study"][0]["key_knowledge"][1]["text"].startswith(
        "proactive and reactive")
