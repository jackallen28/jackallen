"""A subject that exists but has nothing to teach against.

Reported from a real run: Business Management indexed a few exam-style
questions and then said it had mapped them to "0 of 0" study design points.
Zero as the *denominator* means the design itself was empty — so every
question was filed against nothing, and the run reported that as success.

The cause was the importer: it warned "no key knowledge bullets found" and
then wrote the file anyway.
"""

import pytest
import yaml

from blitz.ingest.extract import SourceError
from blitz.studydesign.importer import EmptyStudyDesign, import_study_design


def _course_outline(path):
    """Unit headings and prose, but no dot points — a course outline, not a
    study design. This is the kind of file that produced the empty import."""
    from docx import Document

    d = Document()
    d.add_heading("VCE Legal Studies: Course Outline 2026", 0)
    d.add_heading("Unit 3: Rights and justice", 1)
    d.add_paragraph("In this unit students examine the methods and institutions "
                    "in the justice system. Assessment is two SACs and an exam.")
    d.add_heading("Unit 4: The people and the law", 1)
    d.add_paragraph("Students explore the relationship between the Australian "
                    "people, the Constitution and law-making bodies.")
    d.save(str(path))
    return path


class TestTheImporter:
    def test_a_design_with_no_dot_points_is_refused(self, tmp_path):
        with pytest.raises(EmptyStudyDesign) as exc:
            import_study_design(_course_outline(tmp_path / "outline.docx"),
                                "legal-studies", out_dir=tmp_path)
        assert "No key knowledge dot points" in str(exc.value)

    def test_nothing_is_written_when_it_is_refused(self, tmp_path):
        with pytest.raises(EmptyStudyDesign):
            import_study_design(_course_outline(tmp_path / "outline.docx"),
                                "legal-studies", out_dir=tmp_path)
        assert not (tmp_path / "legal-studies.yaml").exists(), (
            "a subject that cannot be taught against must not be installed")

    def test_the_message_says_where_to_get_the_real_thing(self, tmp_path):
        with pytest.raises(EmptyStudyDesign) as exc:
            import_study_design(_course_outline(tmp_path / "outline.docx"),
                                "legal-studies", out_dir=tmp_path)
        assert "vcaa.vic.edu.au" in str(exc.value)

    def test_the_real_study_design_still_imports(self, tmp_path):
        """The guard must not block the file it is meant to accept."""
        from pathlib import Path

        src = Path(__file__).resolve().parent.parent / "sources" / "vcaa"
        design = src / "business-management-sd.docx"
        if not design.exists():
            pytest.skip("the VCAA source file is not in this checkout")
        out = import_study_design(design, "business-management", out_dir=tmp_path)
        data = yaml.safe_load(out.read_text(encoding="utf-8"))
        n = sum(len(a["key_knowledge"])
                for u in data["units"] for a in u["areas_of_study"])
        assert n == 46


def _empty_design(conn, tmp_path, monkeypatch):
    """Install a subject whose design parses but holds no dot points."""
    from blitz import config
    from blitz.studydesign import loader

    doc = {
        "subject_id": "hollow", "subject_name": "Hollow Subject",
        "units": [{"id": "hollow-u3", "number": 3, "title": "Unit 3",
                   "areas_of_study": [{"id": "hollow-u3-aos1", "number": 1,
                                       "title": "AOS 1", "outcome": "",
                                       "key_knowledge": [], "key_skills": []}]}],
        "question_types": [{"id": "hollow-short", "label": "Short answer",
                            "description": "x"}],
        "command_terms": {},
    }
    d = tmp_path / "designs"
    d.mkdir(exist_ok=True)
    (d / "hollow.yaml").write_text(yaml.safe_dump(doc), encoding="utf-8")
    monkeypatch.setattr(config, "USER_DESIGN_DIR", d)
    monkeypatch.setattr(loader, "USER_DESIGN_DIR", d)
    loader.load_study_design.cache_clear()
    return d


class TestIndexing:
    def test_indexing_against_an_empty_design_is_refused(
            self, conn, tmp_path, monkeypatch):
        """This is the "0 of 0" the report described, caught before it runs."""
        import sys

        sys.path.insert(0, "tests")
        from publisher_pages import build

        from blitz.ingest.index import index_book

        _empty_design(conn, tmp_path, monkeypatch)
        book, _ = build("practice-sac", tmp_path)
        with pytest.raises(SourceError) as exc:
            index_book(conn, book, subject_id="hollow", source_id="s",
                       progress=lambda s: None)
        assert "no dot points" in str(exc.value)
        assert "re-import the study design" in str(exc.value)

    def test_nothing_is_indexed_when_it_is_refused(
            self, conn, tmp_path, monkeypatch):
        import sys

        sys.path.insert(0, "tests")
        from publisher_pages import build

        from blitz.ingest.index import index_book

        _empty_design(conn, tmp_path, monkeypatch)
        book, _ = build("practice-sac", tmp_path)
        with pytest.raises(SourceError):
            index_book(conn, book, subject_id="hollow", source_id="s",
                       progress=lambda s: None)
        assert conn.execute("SELECT COUNT(*) n FROM question").fetchone()["n"] == 0


class TestLexiconAdoption:
    """A self-imported subject should get the lexicon that ships for it.

    Someone who imports Business Management from VCAA themselves was getting
    word-overlap tagging while a lexicon covering all 46 of its dot points sat
    unused in the download.
    """

    def _import_bm(self, tmp_path, monkeypatch):
        from pathlib import Path

        from blitz import config, setup
        from blitz.studydesign import loader

        src = Path(__file__).resolve().parent.parent / "sources" / "vcaa"
        if not (src / "business-management-sd.docx").exists():
            pytest.skip("the VCAA source file is not in this checkout")

        root = tmp_path / "root"
        (root / "study-designs").mkdir(parents=True)
        monkeypatch.setattr(config, "ROOT", root)
        monkeypatch.setattr(config, "USER_DESIGN_DIR", root / "study-designs")
        monkeypatch.setattr(loader, "USER_DESIGN_DIR", root / "study-designs")
        loader.load_study_design.cache_clear()
        return setup.add_study_design(
            src / "business-management-sd.docx", "Business Management", root=root)

    def test_the_shipped_lexicon_is_adopted(self, tmp_path, monkeypatch):
        result = self._import_bm(tmp_path, monkeypatch)
        assert result["lexicon"] == 46

    def test_the_lexicon_file_lands_beside_the_design(self, tmp_path, monkeypatch):
        self._import_bm(tmp_path, monkeypatch)
        designs = tmp_path / "root" / "study-designs"
        assert (designs / "business-management-lexicon.yaml").exists()

    def test_a_lexicon_that_does_not_fit_is_not_adopted(self, tmp_path, monkeypatch):
        """Dot point ids are positional. A lexicon that points at the wrong
        concepts is worse than none, so it is only taken when every id fits."""
        from blitz import config, setup
        from blitz.studydesign import loader

        root = tmp_path / "root"
        (root / "study-designs").mkdir(parents=True)
        monkeypatch.setattr(config, "ROOT", root)
        monkeypatch.setattr(config, "USER_DESIGN_DIR", root / "study-designs")
        monkeypatch.setattr(loader, "USER_DESIGN_DIR", root / "study-designs")

        # A design under the physics name whose ids are nothing like physics'.
        doc = {
            "subject_id": "physics", "subject_name": "Physics",
            "units": [{"id": "px-u3", "number": 3, "title": "U3",
                       "areas_of_study": [{
                           "id": "px-u3-aos1", "number": 1, "title": "A",
                           "outcome": "", "key_skills": [],
                           "key_knowledge": [{"id": "px-u3-aos1-kk01",
                                              "text": "something", "verified": True}]}]}],
            "question_types": [{"id": "q", "label": "Q", "description": "d"}],
            "command_terms": {},
        }
        (root / "study-designs" / "physics.yaml").write_text(
            yaml.safe_dump(doc), encoding="utf-8")
        loader.load_study_design.cache_clear()

        assert setup.adopt_shipped_lexicon("physics", root=root) == 0
        assert not (root / "study-designs" / "physics-lexicon.yaml").exists()


class TestWrongSubject:
    """Picking the wrong subject in the dropdown is easy, and the result is
    worse than nothing: the few questions that match something get filed
    under it."""

    def _report(self, found, untagged, subject="VCE Physics"):
        from blitz.ingest.index import IndexReport

        return IndexReport(source_id="s", subject_id="physics",
                           subject_name=subject, questions_found=found,
                           questions_untagged=untagged)

    def test_a_book_that_mostly_matched_nothing_is_flagged(self):
        assert self._report(10, 8).wrong_subject_suspected

    def test_a_book_that_mostly_matched_is_not(self):
        assert not self._report(10, 2).wrong_subject_suspected

    def test_a_tiny_extract_is_not_flagged(self):
        """Two questions from a page proves nothing either way."""
        assert not self._report(2, 2).wrong_subject_suspected

    def test_the_warning_names_the_subject_and_says_what_to_do(self):
        text = self._report(10, 9).summary()
        assert "VCE Physics" in text
        assert "index it under the right subject" in text

    def test_a_good_index_carries_no_warning(self):
        assert "matched no dot point" not in self._report(10, 1).summary()
