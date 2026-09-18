"""Students and classes: who has been given which question.

Every sheet made for a named student or class is recorded, and two kinds of
workbook are rewritten from the record each time:

    students/MASTER.xlsx        everything, across every student and class
    students/<Name>.xlsx        one per student or class

They are plain Excel files, so they open anywhere and can be sorted and
filtered, and because they are regenerated from the database on every change
they never drift from what the tool actually did. Edit the database through
the app; treat the workbooks as the printout.
"""

from __future__ import annotations

import re
import sqlite3
from datetime import datetime
from pathlib import Path

from .config import STUDENTS_DIR
from .studydesign import load_study_design

LOG_HEAD = ["Date", "Student / class", "Subject", "Sheet", "Question", "Serial",
            "Source", "Dot points", "Type", "Marks", "Sheet file"]


def _label(designs: dict, subject_id: str, kk_id: str) -> str:
    try:
        design = designs.get(subject_id) or load_study_design(subject_id)
        designs[subject_id] = design
        return design.key_knowledge(kk_id).display
    except Exception:
        return kk_id


def log_rows(conn: sqlite3.Connection, student_id: str | None = None) -> list[list]:
    """One row per question given, newest sheet first."""
    designs: dict = {}
    where = "WHERE s.student_id IS NOT NULL"
    params: list = []
    if student_id:
        where += " AND s.student_id = ?"
        params.append(student_id)
    rows = conn.execute(
        f"SELECT s.id AS sheet_id, s.created_at, s.title, s.subject_id, s.pdf_path, "
        f"st.name AS student, sq.position, q.serial, q.citation, q.question_type, "
        f"q.marks, q.id AS qid "
        f"FROM sheet s JOIN student st ON st.id = s.student_id "
        f"JOIN sheet_question sq ON sq.sheet_id = s.id "
        f"JOIN question q ON q.id = sq.question_id "
        f"{where} ORDER BY s.created_at DESC, sq.position", params).fetchall()
    out = []
    for r in rows:
        kks = [x["kk_id"] for x in conn.execute(
            "SELECT kk_id FROM question_kk WHERE question_id = ? ORDER BY primary_kk DESC",
            (r["qid"],))]
        out.append([
            r["created_at"][:16].replace("T", " "),
            r["student"], r["subject_id"], r["title"] or r["sheet_id"],
            r["position"], r["serial"], r["citation"],
            "; ".join(_label(designs, r["subject_id"], k) for k in kks),
            r["question_type"], r["marks"],
            Path(r["pdf_path"]).name if r["pdf_path"] else "",
        ])
    return out


def summary_rows(conn: sqlite3.Connection) -> list[list]:
    """One row per student or class: sheets, questions, last sheet."""
    rows = conn.execute(
        "SELECT st.id, st.name, COUNT(DISTINCT s.id) AS sheets, "
        "COUNT(sq.question_id) AS questions, MAX(s.created_at) AS last "
        "FROM student st LEFT JOIN sheet s ON s.student_id = st.id "
        "LEFT JOIN sheet_question sq ON sq.sheet_id = s.id "
        "GROUP BY st.id ORDER BY st.name").fetchall()
    return [[r["name"], r["sheets"], r["questions"],
             (r["last"] or "")[:16].replace("T", " ")] for r in rows]


def question_rows(conn: sqlite3.Connection) -> list[list]:
    """One row per question ever given: how often, to whom."""
    rows = conn.execute(
        "SELECT q.serial, q.citation, q.subject_id, COUNT(*) AS times, "
        "GROUP_CONCAT(DISTINCT st.name) AS who "
        "FROM sheet_question sq JOIN sheet s ON s.id = sq.sheet_id "
        "JOIN student st ON st.id = s.student_id "
        "JOIN question q ON q.id = sq.question_id "
        "GROUP BY q.id ORDER BY times DESC, q.serial").fetchall()
    return [[r["serial"], r["citation"], r["subject_id"], r["times"], r["who"]]
            for r in rows]


def _autosize(ws) -> None:
    from openpyxl.utils import get_column_letter

    for i, col in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)


def _sheet(wb, title: str, head: list[str], rows: list[list]):
    from openpyxl.styles import Font, PatternFill

    ws = wb.create_sheet(title)
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="E8792B")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    _autosize(ws)
    return ws


def safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9 _-]+", "", name).strip() or "student"


def write_workbooks(conn: sqlite3.Connection, out_dir: Path | None = None) -> list[Path]:
    """Rewrite MASTER.xlsx and one workbook per student. Returns what was written."""
    from openpyxl import Workbook

    out_dir = Path(out_dir or STUDENTS_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")

    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Students", ["Student / class", "Sheets", "Questions", "Last sheet"],
           summary_rows(conn))
    _sheet(wb, "Log", LOG_HEAD, log_rows(conn))
    _sheet(wb, "Questions", ["Serial", "Source", "Subject", "Times given", "Given to"],
           question_rows(conn))
    about = wb.create_sheet("About")
    about.append(["Blitz master record"])
    about.append([f"Rewritten {stamp} from the app's database. Edit nothing here; "
                  "it is regenerated after every sheet."])
    master = out_dir / "MASTER.xlsx"
    wb.save(master)
    written.append(master)

    for st in conn.execute("SELECT id, name FROM student ORDER BY name"):
        wb = Workbook()
        wb.remove(wb.active)
        rows = log_rows(conn, st["id"])
        _sheet(wb, "Questions given", LOG_HEAD, rows)
        sheets = conn.execute(
            "SELECT s.created_at, s.title, s.subject_id, s.pdf_path, "
            "COUNT(sq.question_id) AS n FROM sheet s "
            "LEFT JOIN sheet_question sq ON sq.sheet_id = s.id "
            "WHERE s.student_id = ? GROUP BY s.id ORDER BY s.created_at DESC",
            (st["id"],)).fetchall()
        _sheet(wb, "Sheets", ["Date", "Sheet", "Subject", "Questions", "File"],
               [[r["created_at"][:16].replace("T", " "), r["title"], r["subject_id"],
                 r["n"], Path(r["pdf_path"]).name if r["pdf_path"] else ""]
                for r in sheets])
        path = out_dir / f"{safe_name(st['name'])}.xlsx"
        wb.save(path)
        written.append(path)
    return written
